# -*- coding: utf-8 -*-
"""生成一份用于端到端测试的样例课件：八年级物理《牛顿第一定律与惯性》"""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
layout_title = prs.slide_layouts[0]
layout_bullets = prs.slide_layouts[1]

def add_bullets(slide, items):
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, (text, level) in enumerate(items):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level

# 1 封面
s = prs.slides.add_slide(layout_title)
s.shapes.title.text = "牛顿第一定律与惯性"
s.placeholders[1].text = "八年级物理 · 第八章第2节"

# 2 情境导入
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "情境导入：力与运动的关系"
add_bullets(s, [
    ("思考：正在运动的物体，如果撤去所有外力，它会怎样？", 0),
    ("亚里士多德：必须有力作用在物体上，物体才能运动", 1),
    ("伽利略：运动不需要力来维持，力是改变运动状态的原因", 1),
    ("谁对谁错？让我们用实验说话", 0),
])
s.notes_slide.notes_text_frame.text = "本页通过历史争论引入，强调科学史背景，引导学生质疑直觉。"

# 3 阻力实验
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "探究实验：阻力对运动的影响"
add_bullets(s, [
    ("实验装置：同一斜面、同一高度释放小车，水平面粗糙程度不同", 0),
    ("毛巾表面：小车滑行距离最短，速度减小得最快", 1),
    ("棉布表面：滑行距离较长", 1),
    ("木板表面：滑行距离最长，速度减小得最慢", 1),
    ("结论：平面越光滑，小车受到的阻力越小，滑行得越远", 0),
])
s.notes_slide.notes_text_frame.text = "控制变量法：强调同一小车、同一高度、只改变接触面粗糙程度。"

# 4 牛顿第一定律
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "牛顿第一定律"
add_bullets(s, [
    ("内容：一切物体在没有受到力的作用时，总保持静止状态或匀速直线运动状态", 0),
    ("建立方法：实验 + 科学推理（理想实验法）", 1),
    ("它不能直接用实验验证，是在实验基础上合理外推得到的", 1),
    ("意义：揭示了力与运动的正确关系——力不是维持运动的原因", 0),
])
s.notes_slide.notes_text_frame.text = "重点：定律的两个状态——静止或匀速直线运动。科学推理法是重要思想方法。"

# 5 惯性
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "惯性：物体的固有属性"
add_bullets(s, [
    ("定义：物体保持原来运动状态不变的性质，叫惯性", 0),
    ("一切物体在任何情况下都具有惯性，惯性是物体的固有属性", 1),
    ("惯性大小只与质量有关：质量越大，惯性越大", 1),
    ("与速度无关！速度大不代表惯性大", 1),
    ("生活现象：急刹车时身体前倾、拍打衣服除尘、跳远前助跑", 0),
])

# 6 易错辨析
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "易错辨析"
add_bullets(s, [
    ("辨析1：高速飞行的子弹惯性大？——错，惯性只与质量有关", 0),
    ("辨析2：静止的物体没有惯性？——错，一切物体都有惯性", 0),
    ("辨析3：汽车刹车时人向前倾是受到向前的力？——错，是惯性表现，不存在这个力", 0),
    ("注意：惯性是属性不是力，不能说“受到惯性作用”，要说“由于惯性”", 1),
])
s.notes_slide.notes_text_frame.text = "高频考点：惯性与力的区别，语言表述规范。"

# 7 课堂小结
s = prs.slides.add_slide(layout_bullets)
s.shapes.title.text = "课堂小结与作业"
add_bullets(s, [
    ("伽利略理想实验 → 牛顿第一定律（实验+推理）", 0),
    ("力是改变物体运动状态的原因，不是维持运动的原因", 1),
    ("惯性：固有属性，只与质量有关", 1),
    ("作业：课本P42 1-4题；观察生活中三个惯性现象并解释", 0),
])

out = r"C:/Users/unfurlyang/WorkBuddy/2026-10-06-16-31-56/ppt-study-hub/tests/样例课件-牛顿第一定律.pptx"
import os
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("saved:", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
